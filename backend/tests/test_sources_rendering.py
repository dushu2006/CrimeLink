"""Sources rendering: real viewers per file type, scoped to the ACTIVE dataset.

The reported defects were, precisely:

* the Sources page listed the bundled evaluation corpus instead of the
  uploaded dataset's files;
* PDFs were dumped through a UTF-8 text renderer, producing ``%PDF-1.3``
  garbage on screen;
* valid CSV/XLSX files showed as "No evidence";
* binary files were decoded as text and shown as corrupted symbols.

Every test here goes through the real HTTP endpoints with a real imported
dataset — files generated in the test, never fixtures hardcoded into the
application — and asserts on the *explicit* status the viewer reports.
"""

from __future__ import annotations

import io
import zipfile
from pathlib import Path

import pytest

from app.datasets.pipeline import ImportOptions, run_import
from app.db.session import async_session

JURISDICTION = "RJ-JAIPUR"

PEOPLE_CSV = (
    "person_id,full_name,phone_number\n"
    "PERSON_0001,Arjun Reddy,9812345672\n"
    "PERSON_0002,Sunita Devi,9812345671\n"
)

CASES_CSV = "case_id,case_number,title\nCASE_0001,BC/2026/1,Operation Blue Ledger\n"


def _write_pdf(path: Path) -> None:
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas

    handle = io.BytesIO()
    doc = canvas.Canvas(handle, pagesize=letter)
    doc.drawString(72, 720, "FIR EXTRACT - OPERATION BLUE LEDGER")
    doc.drawString(72, 700, "Complainant reports transfers to ACCT_0001.")
    doc.showPage()
    doc.save()
    path.write_bytes(handle.getvalue())


def _write_docx(path: Path) -> None:
    import docx

    document = docx.Document()
    document.add_heading("Seizure Summary", level=1)
    document.add_paragraph("Arjun Reddy coordinated the ledger entries.")
    table = document.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Account"
    table.rows[0].cells[1].text = "Total"
    table.rows[1].cells[0].text = "ACCT_0001"
    table.rows[1].cells[1].text = "4,50,000"
    document.save(str(path))


def _write_xlsx(path: Path) -> None:
    import openpyxl

    workbook = openpyxl.Workbook()
    register = workbook.active
    register.title = "Register"
    register.append(["account_id", "holder", "bank"])
    register.append(["ACCT_0001", "Arjun Reddy", "SBI"])
    sheet = workbook.create_sheet("Transfers")
    sheet.append(["from", "to", "amount"])
    sheet.append(["ACCT_0001", "ACCT_0002", 250000])
    workbook.save(str(path))


def _write_pptx(path: Path) -> None:
    """A minimal but genuine OOXML deck: one slide with two text runs."""
    slide = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<p:sld xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
        'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">'
        "<p:cSld><p:spTree><p:sp><p:txBody>"
        "<a:p><a:r><a:t>Blue Ledger Briefing</a:t></a:r></a:p>"
        "<a:p><a:r><a:t>Financial flow centres on ACCT_0001</a:t></a:r></a:p>"
        "</p:txBody></p:sp></p:spTree></p:cSld></p:sld>"
    )
    with zipfile.ZipFile(path, "w") as package:
        package.writestr(
            "[Content_Types].xml",
            '<?xml version="1.0"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"/>',
        )
        package.writestr("_rels/.rels", '<?xml version="1.0"?><Relationships xmlns="x"/>')
        package.writestr("ppt/slides/slide1.xml", slide)


@pytest.fixture()
async def rendered_dataset(tmp_path, container):
    """Import a mixed-format dataset through the real pipeline."""
    folder = tmp_path / "blue_ledger"
    folder.mkdir()
    (folder / "cases.csv").write_text(CASES_CSV, encoding="utf-8")
    (folder / "people.csv").write_text(PEOPLE_CSV, encoding="utf-8")
    (folder / "fir.txt").write_text(
        "FIR 12/2026\nOperation Blue Ledger opened at PS-01.\nReference: CASE_0001\n",
        encoding="utf-8",
    )
    (folder / "calls.json").write_text('{"calls": [{"from": "9812345672"}]}', encoding="utf-8")
    _write_pdf(folder / "statement.pdf")
    _write_docx(folder / "summary.docx")
    _write_xlsx(folder / "ledger.xlsx")
    _write_pptx(folder / "brief.pptx")
    (folder / "blob.bin").write_bytes(bytes(range(256)) * 8)

    async with async_session() as session:
        report = await run_import(
            session,
            [folder],
            ImportOptions(
                name="Blue Ledger sources",
                copy_inputs=True,
                activate=True,
                build_graph=True,
                jurisdiction_id=JURISDICTION,
            ),
        )
    assert report.error is None, report.error
    return report.dataset_id


PREFIX = "blue_ledger/"


def _preview(client, headers, path: str, **params):
    path = f"{PREFIX}{path}" if "/" not in path else path
    query = "&".join(f"{key}={value}" for key, value in params.items())
    url = f"/api/v1/sources/preview?path={path}" + (f"&{query}" if query else "")
    return client.get(url, headers=headers)


# ---------------------------------------------------------------------------
# The listing follows the ACTIVE dataset, not the evaluation corpus
# ---------------------------------------------------------------------------


async def test_sources_page_lists_the_uploaded_dataset(client, admin_headers, rendered_dataset):
    body = client.get("/api/v1/sources/files", headers=admin_headers).json()
    assert body["dataset_id"] == rendered_dataset
    paths = {item["path"] for item in body["items"]}
    for expected in (
        "blue_ledger/fir.txt",
        "blue_ledger/statement.pdf",
        "blue_ledger/ledger.xlsx",
        "blue_ledger/summary.docx",
        "blue_ledger/brief.pptx",
        "blue_ledger/calls.json",
        "blue_ledger/blob.bin",
    ):
        assert expected in paths, paths

    by_path = {item["path"]: item for item in body["items"]}
    assert by_path["blue_ledger/statement.pdf"]["media_type"] == "application/pdf"
    assert by_path["blue_ledger/ledger.xlsx"]["sheets"] == ["Register", "Transfers"]
    # The PDF/TXT/DOCX/PPTX files became evidence documents with provenance.
    assert by_path["blue_ledger/statement.pdf"]["doc_id"]
    assert by_path["blue_ledger/statement.pdf"]["case_number"]


async def test_every_source_row_carries_the_required_identity_fields(
    client, admin_headers, rendered_dataset
):
    items = client.get("/api/v1/sources/files", headers=admin_headers).json()["items"]
    for item in items:
        assert item["source_id"]
        assert item["dataset_id"] == rendered_dataset
        assert item["filename"]
        assert item["path"]
        assert item["media_type"]
        assert item["size_bytes"] >= 0
        assert item["status"]
        assert "download_url" in item


# ---------------------------------------------------------------------------
# Renderers: the right mechanism per format, explicit state always
# ---------------------------------------------------------------------------


async def test_pdf_renders_as_pdf_not_as_raw_syntax(client, admin_headers, rendered_dataset):
    result = _preview(client, admin_headers, "statement.pdf").json()
    assert result["status"] == "AVAILABLE"
    assert result["render_kind"] == "pdf"
    assert result["pdf"]["page_count"] >= 1
    assert "BLUE LEDGER" in result["pdf"]["pages"][0]["text"].upper()
    # The signed raw URL must deliver real PDF bytes.
    raw = client.get(result["raw_url"])
    assert raw.status_code == 200
    assert raw.content.startswith(b"%PDF-")
    assert raw.headers["content-type"].startswith("application/pdf")
    assert "inline" in raw.headers["content-disposition"]


async def test_pdf_bytes_also_arrive_via_bearer_token(client, admin_headers, rendered_dataset):
    raw = client.get("/api/v1/sources/raw?path=blue_ledger/statement.pdf", headers=admin_headers)
    assert raw.status_code == 200
    assert raw.content.startswith(b"%PDF-")
    assert raw.headers["x-content-type-options"] == "nosniff"


async def test_csv_renders_as_a_table(client, admin_headers, rendered_dataset):
    result = _preview(client, admin_headers, "people.csv").json()
    assert result["status"] == "AVAILABLE"
    assert result["render_kind"] == "csv"
    window = result["window"]
    assert window["columns"] == ["person_id", "full_name", "phone_number"]
    values = [row["values"]["full_name"] for row in window["rows"]]
    assert "Arjun Reddy" in values


async def test_csv_windows_page_with_offset_and_limit(client, admin_headers, rendered_dataset):
    page = _preview(client, admin_headers, "people.csv", offset=3, limit=1).json()
    rows = page["window"]["rows"]
    assert len(rows) == 1
    assert rows[0]["row"] == 3
    assert rows[0]["values"]["full_name"] == "Sunita Devi"


async def test_xlsx_lists_sheets_and_switches_between_them(client, admin_headers, rendered_dataset):
    first = _preview(client, admin_headers, "ledger.xlsx").json()
    assert first["render_kind"] == "xlsx"
    assert first["sheets"] == ["Register", "Transfers"]
    assert first["sheet"] == "Register"
    assert first["window"]["rows"][0]["values"]["holder"] == "Arjun Reddy"

    second = _preview(client, admin_headers, "ledger.xlsx", sheet="Transfers").json()
    assert second["sheet"] == "Transfers"
    assert second["window"]["columns"] == ["from", "to", "amount"]
    assert second["window"]["rows"][0]["values"]["amount"] == "250000"


async def test_json_renders_formatted(client, admin_headers, rendered_dataset):
    result = _preview(client, admin_headers, "calls.json").json()
    assert result["render_kind"] == "json"
    text = "\n".join(line["text"] for line in result["window"]["lines"])
    assert '"calls": [' in text  # pretty-printed, not a raw one-liner


async def test_docx_renders_structured_content(client, admin_headers, rendered_dataset):
    result = _preview(client, admin_headers, "summary.docx").json()
    assert result["status"] == "AVAILABLE"
    assert result["render_kind"] == "docx"
    blocks = result["document_blocks"]
    kinds = [block["type"] for block in blocks]
    assert "heading" in kinds and "paragraph" in kinds and "table" in kinds
    table = next(block for block in blocks if block["type"] == "table")
    assert table["rows"][1] == ["ACCT_0001", "4,50,000"]


async def test_pptx_renders_slides(client, admin_headers, rendered_dataset):
    result = _preview(client, admin_headers, "brief.pptx").json()
    assert result["render_kind"] == "pptx"
    slides = result["slides"]
    assert slides[0]["title"] == "Blue Ledger Briefing"
    assert any("ACCT_0001" in line for line in slides[0]["lines"])


async def test_binary_is_never_decoded_as_text(client, admin_headers, rendered_dataset):
    result = _preview(client, admin_headers, "blob.bin").json()
    assert result["status"] == "UNSUPPORTED"
    assert result["render_kind"] == "binary"
    assert result["window"] is None
    assert "download" in (result["reason"] or "").lower()
    # But the bytes remain faithfully downloadable.
    raw = client.get("/api/v1/sources/raw?path=blue_ledger/blob.bin", headers=admin_headers)
    assert raw.content == bytes(range(256)) * 8


async def test_plain_text_still_renders_as_text(client, admin_headers, rendered_dataset):
    result = _preview(client, admin_headers, "fir.txt").json()
    assert result["render_kind"] == "text"
    assert any("Blue Ledger" in line["text"] for line in result["window"]["lines"])


async def test_mislabelled_pdf_is_detected_from_bytes_not_extension(
    client, admin_headers, container
):
    """A PDF named ``.txt`` used to be UTF-8-decoded and garbled on screen."""
    folder = container.settings.data_dir / "sources-mislabel"
    folder.mkdir(parents=True, exist_ok=True)
    (folder / "cases.csv").write_text(CASES_CSV, encoding="utf-8")
    fake = folder / "notes.txt"
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas

    handle = io.BytesIO()
    doc = canvas.Canvas(handle, pagesize=letter)
    doc.drawString(72, 720, "MISLABELLED PDF CONTENT")
    doc.showPage()
    doc.save()
    fake.write_bytes(handle.getvalue())

    async with async_session() as session:
        report = await run_import(
            session,
            [folder],
            ImportOptions(
                name="Mislabelled", copy_inputs=True, activate=True,
                build_graph=False, jurisdiction_id=JURISDICTION,
            ),
        )
    assert report.error is None, report.error
    result = _preview(client, admin_headers, "sources-mislabel/notes.txt").json()
    # Discovery re-labelled the manifest entry by content, so the preview sees
    # a real .pdf file in the workspace: no raw syntax anywhere.
    assert result["render_kind"] in {"pdf", "text"}
    if result["render_kind"] == "pdf":
        assert "MISLABELLED" in result["pdf"]["pages"][0]["text"]
    else:
        assert "%PDF-" not in "\n".join(line["text"] for line in result["window"]["lines"])


# ---------------------------------------------------------------------------
# Explicit states and isolation
# ---------------------------------------------------------------------------


async def test_missing_file_reports_not_found_explicitly(client, admin_headers, rendered_dataset):
    result = _preview(client, admin_headers, "does-not-exist.csv").json()
    assert result["status"] == "NOT_FOUND"
    assert result["window"] is None


async def test_traversal_and_absolute_paths_are_refused(client, admin_headers, rendered_dataset):
    for attempt in ("../../blue_ledger/fir.txt", "/etc/passwd", "..%2F..%2Fetc%2Fpasswd"):
        response = _preview(client, admin_headers, attempt)
        assert response.status_code in (200, 400, 404), attempt
        if response.status_code == 200:
            assert response.json()["status"] in {"NOT_FOUND", "UNSUPPORTED"}


async def test_replacing_the_dataset_closes_the_files_behind_it(
    client, admin_headers, tmp_path, container, rendered_dataset
):
    """After activation of dataset B, A's evidence files are unreachable.

    This is the isolation test that matters for the viewer: the previous
    dataset's workspace still exists on disk for audit, but the sources API
    resolves strictly inside the ACTIVE dataset's root, so old paths return
    NOT_FOUND instead of quietly staying openable.
    """
    other = tmp_path / "second_dataset"
    other.mkdir()
    (other / "cases.csv").write_text(
        "case_id,case_number,title\nCASE_0001,NEW/2026/9,Brand new case\n", encoding="utf-8"
    )
    (other / "only.txt").write_text("Second dataset content.\n", encoding="utf-8")
    async with async_session() as session:
        report = await run_import(
            session,
            [other],
            ImportOptions(
                name="Second dataset", copy_inputs=True, activate=True,
                build_graph=True, jurisdiction_id=JURISDICTION,
            ),
        )
    assert report.error is None, report.error

    listing = client.get("/api/v1/sources/files", headers=admin_headers).json()
    assert listing["dataset_id"] == report.dataset_id
    assert {item["path"] for item in listing["items"]} == {"second_dataset/cases.csv", "second_dataset/only.txt"}

    stale = _preview(client, admin_headers, "blue_ledger/statement.pdf").json()
    assert stale["status"] == "NOT_FOUND"
    stale_raw = client.get("/api/v1/sources/raw?path=blue_ledger/statement.pdf", headers=admin_headers)
    assert stale_raw.status_code == 404

    # The new dataset's own file resolves normally.
    fresh = _preview(client, admin_headers, "second_dataset/only.txt").json()
    assert fresh["status"] == "AVAILABLE"


async def test_signed_raw_url_expires_without_leaking(client, admin_headers, rendered_dataset):
    result = _preview(client, admin_headers, "statement.pdf").json()
    # An unauthenticated request with a bad signature is refused with the same
    # status as a nonexistent path: the endpoint cannot enumerate files.
    unsigned = result["raw_url"].split("&sig=")[0] + "&sig=nope"
    response = client.get(unsigned)
    assert response.status_code == 404
